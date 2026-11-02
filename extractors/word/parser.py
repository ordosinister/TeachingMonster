import json
import os
import uuid
from pathlib import Path

from docx import Document
from langchain_core.messages import HumanMessage

from extractors.shared.utils import image_to_base64
from src.path.path_definition import get_project_dir


def _extract_images_from_docx(docx_path, output_dir):
    r"""Extract inline images from a DOCX file and save them to output_dir.

    Returns a dict mapping (paragraph_idx, run_idx) to image metadata dict.
    """
    import zipfile
    from xml.etree import ElementTree as ET

    image_dir = output_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)

    image_map = {}
    rels_map = {}

    with zipfile.ZipFile(docx_path, 'r') as z:
        # Build relationship map: rId to target (image path inside docx)
        try:
            rels_xml = z.read('word/_rels/document.xml.rels')
            root = ET.fromstring(rels_xml)
            ns = {'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}
            for rel in root.findall('rel:Relationship', ns):
                rId = rel.get('Id')
                target = rel.get('Target')
                rel_type = rel.get('Type', '')
                if 'image' in rel_type:
                    rels_map[rId] = target
        except KeyError:
            pass

        # Parse document.xml to locate image placements
        try:
            doc_xml = z.read('word/document.xml')
            root = ET.fromstring(doc_xml)
        except KeyError:
            return image_map

        # WordprocessingML namespace
        w_ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        ns = {'w': w_ns}

        paragraphs = root.findall('.//w:p', ns)
        for para_idx, para in enumerate(paragraphs):
            runs = para.findall('.//w:r', ns)
            for run_idx, run in enumerate(runs):
                # Search for blip in explicit namespace form
                blip_ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
                blips = run.findall(f'.//{{{blip_ns}}}blip')
                if not blips:
                    blips = run.findall('.//a:blip', {'a': blip_ns})
                for blip in blips:
                    embed = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if embed and embed in rels_map:
                        img_path_in_docx = 'word/' + rels_map[embed]
                        try:
                            img_bytes = z.read(img_path_in_docx)
                            ext = Path(img_path_in_docx).suffix.lstrip('.') or 'png'
                            filename = f"section_{para_idx}_run_{run_idx}.{ext}"
                            filepath = image_dir / filename
                            filepath.write_bytes(img_bytes)
                            image_map[(para_idx, run_idx)] = {
                                "path": str(filepath),
                                "filename": filename,
                                "content_type": f"image/{ext}" if ext != 'emf' else "image/png",
                            }
                        except KeyError:
                            pass
    return image_map


def parse_docx(docx_path):
    doc = Document(docx_path)

    output_dir = Path(os.path.join(get_project_dir(), "data", str(uuid.uuid4())))
    output_dir.mkdir(parents=True, exist_ok=True)

    image_map = _extract_images_from_docx(docx_path, output_dir)

    elements = []
    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        style = para.style.name if para.style else "Normal"

        # Check for inline image in this paragraph
        img_meta = None
        for run_idx in range(len(para.runs)):
            if (para_idx, run_idx) in image_map:
                img_meta = image_map[(para_idx, run_idx)]
                break

        if img_meta:
            elements.append({
                "element_id": f"para_{para_idx}",
                "type": "image",
                "style": style,
                "path": img_meta["path"],
                "content_type": img_meta["content_type"],
            })

        if text:
            elements.append({
                "element_id": f"para_{para_idx}",
                "type": "text",
                "style": style,
                "text": text,
            })

    # Extract tables
    for table_idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        elements.append({
            "element_id": f"table_{table_idx}",
            "type": "table",
            "style": "Table",
            "rows": rows,
        })

    doc_data = {
        "document_name": Path(docx_path).name,
        "elements": elements,
    }

    return doc_data, output_dir


def docx_to_base64_messages(docx_data) -> list[HumanMessage]:
    r"""Convert extracted Word document data to Base64 image messages."""
    doc_input = json.dumps(docx_data, ensure_ascii=False)

    messages = [
        HumanMessage(
            content=f"<WORD_INPUT>: {doc_input}"
        )
    ]

    for element in docx_data.get("elements", []):
        if element.get("type") == "image":
            path = element["path"]
            image_str = image_to_base64(path)
            messages.append(
                HumanMessage(
                    content=[
                        {
                            "type": "text",
                            "text": f"Element ID: {element['element_id']}; filepath: {path}"
                        },
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:image/jpeg;base64,{image_str}"},
                        }
                    ]
                )
            )
    return messages


if __name__ == "__main__":
    docx_path = r"data/sample.docx"
    doc_data, output_dir = parse_docx(docx_path)
    print(json.dumps(doc_data, ensure_ascii=False, indent=2))